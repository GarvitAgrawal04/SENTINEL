from pptx import Presentation

prs = Presentation(r"D:\Gta\SENTINEL_SENTINEL2_Deck_EyeComfort_Fixed.pptx")

for i, slide in enumerate(prs.slides):
    print(f"--- Slide {i+1} ---")
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            print(shape.text)
    print("\n")
